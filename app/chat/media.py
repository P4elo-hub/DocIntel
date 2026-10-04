"""Конвертация загруженных медиа в content-part для OpenAI Chat Completions.

Native multimodal: изображения отдаются как `image_url`-part (data: base64
URI) прямо в основной chat.completions.create — без отдельного Vision-вызова.
Голос идёт через Whisper-1 (без FFmpeg, Whisper принимает ogg/m4a/mp3
напрямую). PDF/DOCX — извлечение текста через pypdf / python-docx.

Замечание: для image-part `part["image_url"]["url"]` содержит полный base64
изображения. Если этот dict сохраняется в `ChatMessage.media_refs.part`, он
будет загружен заново при каждом обращении к истории. Для учебной задачи —
это допустимо; продакшен потребовал бы внешнего blob-storage с ID.
"""

import base64
import re
from io import BytesIO

from docx import Document
from fastapi import UploadFile
from openai import AsyncOpenAI
from pypdf import PdfReader

from app.core.exceptions import VoiceUnavailableError

ContentPart = dict  # тип-алиас под OpenAI Chat Completions content-part

VOICE_UNAVAILABLE_MESSAGE = (
    "Сейчас временно не могу обрабатывать голосовые сообщения."
)

_VOICE_PREFIX = "[пользователь сказал голосом]:"
_DOC_PREFIXES = ("[документ PDF]:", "[документ DOCX]:")

# Доменная подсказка для Whisper: перечисляем термины/аббревиатуры «Истории
# операций», которые ASR иначе слышит как частотные слова (OG→ОГЭ, History
# Ops→Хистриопс). Whisper использует это как контекст и точнее распознаёт жаргон.
# Лимит подсказки ~224 токена, поэтому держим короткий список ключевых терминов.
# Переопределяется через LLM_WHISPER_PROMPT (см. app/core/config.py).
WHISPER_DOMAIN_PROMPT = (
    "История операций (History Ops, HistoryOps, HO). "
    "Order Gateway (OG, order-gateway). Composite (Композит), ScreenApi. "
    "GET_LINKED_EVENTS, Linked Events, linkedEvents (не LinkedIn). "
    "Вариационная маржа, вариационной маржи, varmargin_input, varmargin_output. "
    "БПХ (BPH), БК (BC), ДКА (DCA), ОШ, СберИнвестор, СИ2. "
    "gRPC, Thrift, proto, GET_TRADES, GET_DEALS, GET_DETAILS, SEND_DETAILS, "
    "GET_OPERATIONS_WITH_DETAILS, GET_OPERATIONS_FEED, GetTradeOrders. "
    "SIHIST, agreement-gateway, dictionary-gateway, funds_output, TAX_RETAIN."
)


def normalize_domain_query(text: str) -> str:
    """Чинка типичных Whisper-ошибок доменных терминов перед RAG/агентами."""
    t = (text or "").strip()
    if not t:
        return t
    # Whisper: Linked Events → LinkedIn / LinkedIn.com
    t = re.sub(r"(?i)\blinkedin\.com\b", "Linked Events", t)
    t = re.sub(r"(?i)\blinkedin\b", "Linked Events", t)
    t = re.sub(r"(?i)\blinked\s*in\b", "Linked Events", t)
    # Whisper: «GET и OPERATIONS WITH DETAILS» → GET_OPERATIONS_WITH_DETAILS
    t = re.sub(
        r"(?i)\bGET\s+и\s+OPERATIONS\s+WITH\s+DETAILS\b",
        "GET_OPERATIONS_WITH_DETAILS",
        t,
    )
    t = re.sub(
        r"(?i)\bGET\s+OPERATIONS\s+WITH\s+DETAILS\b",
        "GET_OPERATIONS_WITH_DETAILS",
        t,
    )
    # «в рационной моржи/мараже» → «вариационной маржи/марже»
    t = re.sub(
        r"(?i)\bв\s*рационн(\w*)\s+(?:морж|мараж|марж)(\w*)",
        r"вариационн\1 марж\2",
        t,
    )
    t = re.sub(
        r"(?i)\bрационн(\w*)\s+(?:морж|мараж)(\w*)",
        r"вариационн\1 марж\2",
        t,
    )
    t = re.sub(r"(?i)\bмараж", "марж", t)
    return t


# «не про Linked Events» / «не о налогах» — иначе имя чужого API уходит в поиск.
_NEGATED_MENTION_RE = re.compile(
    r"(?i)(?:\bне\s+про\b|\bне\s+о\b|\bне\s+об\b|\bне\s+про\s+этот\b|"
    r"\bбез\b|\bкроме\b|\bnot\s+about\b)\s+"
    r"[«\"'(]?[^.,;:!?\n]{1,80}"
)

# Имя обмена/метода для вырезания отвергнутых упоминаний.
_API_TOKEN_RE = (
    r"(?:Get|Set|Send|Create|Update|Delete|List|Fetch|Post|Put|Patch)"
    r"[A-Z][A-Za-z0-9]+"
    r"|(?:[A-Z][A-Z0-9]+(?:[_-][A-Z0-9]+)+)"
)

# «это SendOperations ты прислал, не тот» / «SendOperations — не то»
_REJECTED_API_SENT_RE = re.compile(
    rf"(?i)(?:это\s+)?({_API_TOKEN_RE})"
    rf"[^.!?\n]{{0,80}}?\bне\s+то[тм]?\b"
)

# «не SendOperations» / «не тот SendOperations»
_REJECTED_NE_API_RE = re.compile(
    rf"(?i)\bне\s+(?:тот\s+|ту\s+|то\s+|та\s+)?({_API_TOKEN_RE})\b"
)


def strip_negated_mentions(text: str) -> str:
    """Убрать отрицательные упоминания и отвергнутые API перед retrieval."""
    cleaned = _NEGATED_MENTION_RE.sub(" ", text or "")
    cleaned = _REJECTED_API_SENT_RE.sub(" ", cleaned)
    cleaned = _REJECTED_NE_API_RE.sub(" ", cleaned)
    return re.sub(r"\s+", " ", cleaned).strip()


def retrieval_query_text(*texts: str) -> str:
    """Текст для RAG: нормализация домена + выкинуть «не про X» / отвергнутый API."""
    blob = "\n".join(t for t in texts if (t or "").strip())
    return strip_negated_mentions(normalize_domain_query(blob))


def extract_rag_query(user_content: str, media_refs: dict | None) -> str:
    """Текст для RAG-поиска: caption/сообщение или транскрипт голоса."""
    text = (user_content or "").strip()
    if not text:
        if not media_refs or not isinstance(media_refs, dict):
            return ""
        part = media_refs.get("part")
        if not part or part.get("type") != "text":
            return ""
        part_text = (part.get("text") or "").strip()
        if part_text.startswith(_VOICE_PREFIX):
            body = part_text[len(_VOICE_PREFIX) :].lstrip("\n")
            text = body.strip()
        else:
            for prefix in _DOC_PREFIXES:
                if part_text.startswith(prefix):
                    body = part_text[len(prefix) :].lstrip("\n")
                    text = body[:2000].strip()
                    break
    return normalize_domain_query(text)


async def media_to_part(
    media: UploadFile, llm_client: AsyncOpenAI
) -> ContentPart:
    """Конвертирует загруженное медиа в content-part для chat.completions.

    Поддерживаемые типы:
    - `image/*` → image_url-part с data: base64 URI.
    - `audio/*` (а также `application/ogg` — Telegram voice) → Whisper-1
      транскрипт → text-part с пометкой '[пользователь сказал голосом]:'.
    - `application/pdf` → pypdf → text-part '[документ PDF]:'.
    - `application/vnd.openxmlformats-officedocument.wordprocessingml.document`
      → python-docx → text-part '[документ DOCX]:'.

    Raises:
        ValueError: для неподдерживаемых MIME-типов.
    """
    mime = media.content_type or ""
    data = await media.read()

    if mime.startswith("image/"):
        b64 = base64.b64encode(data).decode()
        return {
            "type": "image_url",
            "image_url": {"url": f"data:{mime};base64,{b64}"},
        }

    if mime.startswith("audio/") or mime == "application/ogg":
        prompt, language = _whisper_hints()
        transcript = await whisper_transcribe(
            data, media.filename or "audio.ogg", llm_client,
            prompt=prompt, language=language,
        )
        return {
            "type": "text",
            "text": f"[пользователь сказал голосом]:\n{transcript}",
        }

    if mime == "application/pdf":
        return {
            "type": "text",
            "text": f"[документ PDF]:\n{extract_pdf_text(data)[:30_000]}",
        }

    if mime.endswith("wordprocessingml.document"):
        return {
            "type": "text",
            "text": f"[документ DOCX]:\n{extract_docx_text(data)[:30_000]}",
        }

    raise ValueError(f"Unsupported media type: {mime}")


def _whisper_hints() -> tuple[str, str | None]:
    """Подсказка и язык для Whisper из настроек (с фолбэком на встроенный список)."""
    try:
        from app.core.config import get_settings

        llm = get_settings().llm
        prompt = (llm.whisper_prompt or "").strip() or WHISPER_DOMAIN_PROMPT
        language = (llm.whisper_language or "").strip() or None
    except Exception:
        # Настройки недоступны (например, в изолированном тесте) — не роняем голос.
        prompt, language = WHISPER_DOMAIN_PROMPT, None
    return prompt, language


async def whisper_transcribe(
    audio_bytes: bytes,
    filename: str,
    llm_client: AsyncOpenAI,
    *,
    prompt: str | None = None,
    language: str | None = None,
) -> str:
    """Whisper-1 принимает ogg/m4a/mp3/wav/flac/webm без конвертации.

    prompt — доменная подсказка (термины/аббревиатуры), повышает точность ASR на
    жаргоне. language — код ISO-639-1 (напр. "ru"). Пустые значения не передаём.
    """
    f = BytesIO(audio_bytes)
    f.name = filename  # OpenAI SDK ориентируется на расширение из .name
    kwargs: dict = {"model": "whisper-1", "file": f}
    if prompt and prompt.strip():
        kwargs["prompt"] = prompt.strip()
    if language and language.strip():
        kwargs["language"] = language.strip()
    try:
        result = await llm_client.audio.transcriptions.create(**kwargs)
    except VoiceUnavailableError:
        raise
    except Exception as exc:
        raise VoiceUnavailableError(VOICE_UNAVAILABLE_MESSAGE) from exc
    return result.text


def extract_pdf_text(data: bytes, max_pages: int = 50) -> str:
    """Извлекает текст из PDF. Сканы (мало текста на много страниц) — заглушка."""
    reader = PdfReader(BytesIO(data))
    parts: list[str] = []
    for i, page in enumerate(reader.pages):
        if i >= max_pages:
            break
        parts.append(page.extract_text() or "")
    text = "\n\n".join(parts).strip()
    if len(text) < 100 and len(reader.pages) >= 5:
        return "[это скан, OCR пока не поддерживается]"
    return text


def extract_docx_text(data: bytes) -> str:
    """Извлекает текст параграфов и таблиц из DOCX."""
    doc = Document(BytesIO(data))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
